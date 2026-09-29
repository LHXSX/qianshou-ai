#!/usr/bin/env python3
"""sheet_to_text — Excel/CSV 转可读文本 (混合包原子技能)."""
import csv
import io
import json
import os
import sys
import time


def _read_bytes() -> tuple[str, bytes]:
    d = os.environ.get("EC_INPUT_DIR", "")
    if d and os.path.isdir(d):
        for fn in sorted(os.listdir(d)):
            low = fn.lower()
            if low.endswith((".xlsx", ".xls", ".csv")):
                with open(os.path.join(d, fn), "rb") as fh:
                    return fn, fh.read()
    return "stdin", sys.stdin.buffer.read()


def _csv_text(data: bytes) -> str:
    text = data.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    return "\n".join("\t".join(row) for row in reader).strip()


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts).strip()


def main() -> int:
    t0 = time.time()
    name, data = _read_bytes()
    if not data:
        print(json.dumps({"status": "error", "error": "empty_input", "result_text": ""}, ensure_ascii=False))
        return 1
    low = name.lower()
    try:
        if low.endswith(".csv"):
            text = _csv_text(data)
            backend = "csv"
        else:
            try:
                text = _xlsx_text(data)
                backend = "openpyxl"
            except Exception:
                text = _csv_text(data)
                backend = "csv_fallback"
    except Exception as exc:
        print(json.dumps({
            "status": "error",
            "error": str(exc),
            "result_text": "",
        }, ensure_ascii=False))
        return 1
    print(json.dumps({
        "status": "ok",
        "schema_version": "v1",
        "task_type": "sheet_to_text",
        "backend": backend,
        "result_text": text,
        "elapsed_ms": int((time.time() - t0) * 1000),
        "summary": {"chars": len(text)},
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
