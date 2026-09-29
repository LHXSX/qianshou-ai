#!/usr/bin/env python3
"""excel_export — 结构化数据 → Excel(.xlsx) 报告 (企业级 · 2026-06-10)

补平台最致命缺口:此前只能输出 JSON/CSV,客户(律所/HR/财务)要的是
能打开、能筛选、能直接交付的 Excel。本脚本全行业复用。

输入(stdin JSON 或 EC_PARAMS,优先 EC_PARAMS):
  形态 A · 单表:
    {"columns": ["列1","列2"], "rows": [["a","b"], ["c","d"]]}
  形态 B · 字典行(自动推断列):
    {"records": [{"name":"张三","age":30}, ...]}
  形态 C · 多 sheet:
    {"sheets": [{"name":"合同要素","columns":[...],"rows":[...]},
                {"name":"风险清单","records":[...]}]}
  兼容:上游脚本的 {"result_rows": [...]} / {"result_lines": ["{json}", ...]}

参数 (EC_PARAMS):
  filename        str   输出文件名(默认 report.xlsx)
  title           str   首行大标题(可选)
  freeze_header   bool  冻结表头行(默认 true)
  auto_filter     bool  表头加筛选(默认 true)
  max_col_width   int   列宽上限字符(默认 60)
  return_base64   bool  小文件(<4MB)在 stdout 返 base64(默认 true · 无 OSS 也能取)

输出:写 EC_OUTPUT_DIR/<filename> · stdout 返 schema v1 + 元信息(+可选 base64)
"""
import base64
import json
import os
import sys
import time


def _params() -> dict:
    raw = os.environ.get("EC_PARAMS", "")
    if raw:
        try:
            return json.loads(raw) or {}
        except Exception:
            pass
    return {}


def _records_to_table(records: list) -> tuple:
    """[{...}, ...] → (columns, rows) · 列 = 所有键的并集(保序)"""
    columns: list = []
    seen = set()
    for r in records:
        if isinstance(r, dict):
            for k in r.keys():
                if k not in seen:
                    seen.add(k)
                    columns.append(k)
    rows = []
    for r in records:
        if isinstance(r, dict):
            rows.append([r.get(c, "") for c in columns])
        else:
            rows.append([r])
    if not columns:
        columns = ["value"]
    return columns, rows


def _normalize_sheets(data: dict) -> list:
    """把各种输入形态统一成 [{name, columns, rows}, ...]"""
    # 形态 C · 多 sheet
    if isinstance(data.get("sheets"), list) and data["sheets"]:
        out = []
        for i, sh in enumerate(data["sheets"]):
            name = str(sh.get("name") or f"Sheet{i+1}")[:31]
            if isinstance(sh.get("records"), list):
                cols, rows = _records_to_table(sh["records"])
            else:
                cols = sh.get("columns") or []
                rows = sh.get("rows") or []
            out.append({"name": name, "columns": cols, "rows": rows})
        return out

    # 形态 B · records / result_rows
    records = data.get("records") or data.get("result_rows")
    if isinstance(records, list) and records:
        cols, rows = _records_to_table(records)
        return [{"name": str(data.get("sheet_name") or "Sheet1")[:31],
                 "columns": cols, "rows": rows}]

    # 兼容 · result_lines(每行一条 JSON 字符串)
    if isinstance(data.get("result_lines"), list) and data["result_lines"]:
        parsed = []
        for ln in data["result_lines"]:
            try:
                parsed.append(json.loads(ln))
            except Exception:
                parsed.append({"value": ln})
        cols, rows = _records_to_table(parsed)
        return [{"name": "Sheet1", "columns": cols, "rows": rows}]

    # 形态 A · columns + rows
    cols = data.get("columns") or []
    rows = data.get("rows") or []
    if cols or rows:
        return [{"name": str(data.get("sheet_name") or "Sheet1")[:31],
                 "columns": cols, "rows": rows}]

    return []


def _cell(v):
    """xlsx 可写值:基本类型直写,复杂类型转 JSON 字符串"""
    if v is None:
        return ""
    if isinstance(v, (str, int, float, bool)):
        return v
    return json.dumps(v, ensure_ascii=False)


def main() -> int:
    t0 = time.time()
    p = _params()

    # 读输入
    raw = ""
    try:
        raw = sys.stdin.read()
    except Exception:
        raw = ""
    data: dict = {}
    if raw and raw.lstrip().startswith(("{", "[")):
        try:
            obj = json.loads(raw)
            if isinstance(obj, list):
                data = {"records": obj}
            elif isinstance(obj, dict):
                data = obj
        except Exception:
            pass
    # EC_PARAMS 覆盖/补充
    for k in ("sheets", "records", "result_rows", "columns", "rows", "result_lines"):
        if p.get(k) is not None:
            data[k] = p[k]

    sheets = _normalize_sheets(data)
    if not sheets:
        print(json.dumps({
            "status": "failed", "task_type": "excel_export",
            "error": "无可导出数据(需 sheets/records/columns+rows)",
            "summary_text": "❌ 无数据",
        }, ensure_ascii=False))
        return 1

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        print(json.dumps({
            "status": "failed", "task_type": "excel_export",
            "error": f"节点缺 openpyxl: {exc}",
            "summary_text": "❌ pip install openpyxl(应在 lite tier)",
        }, ensure_ascii=False))
        return 1

    filename = p.get("filename") or data.get("filename") or "report.xlsx"
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    title = p.get("title") or data.get("title") or ""
    freeze_header = p.get("freeze_header", True)
    auto_filter = p.get("auto_filter", True)
    max_col_width = int(p.get("max_col_width") or 60)
    return_base64 = p.get("return_base64", True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    total_rows = 0
    for sh in sheets:
        ws = wb.create_sheet(title=(sh["name"] or "Sheet")[:31])
        cols = [str(c) for c in (sh["columns"] or [])]
        rows = sh["rows"] or []
        r0 = 1

        if title:
            ws.cell(row=1, column=1, value=title).font = title_font
            if cols:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
            r0 = 2

        # 表头
        header_row = r0
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=header_row, column=ci, value=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

        # 数据行
        for ri, row in enumerate(rows, header_row + 1):
            for ci in range(1, len(cols) + 1):
                v = row[ci - 1] if ci - 1 < len(row) else ""
                cell = ws.cell(row=ri, column=ci, value=_cell(v))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
            total_rows += 1

        # 列宽自适应
        for ci, c in enumerate(cols, 1):
            maxlen = len(str(c))
            for row in rows[:500]:
                if ci - 1 < len(row):
                    maxlen = max(maxlen, len(str(_cell(row[ci - 1]))))
            ws.column_dimensions[get_column_letter(ci)].width = min(max(10, maxlen + 2), max_col_width)

        if cols:
            if freeze_header:
                ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
            if auto_filter and rows:
                ws.auto_filter.ref = f"{get_column_letter(1)}{header_row}:{get_column_letter(len(cols))}{header_row + len(rows)}"

    # 写盘
    import io
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    output_path = None
    out_dir = os.environ.get("EC_OUTPUT_DIR", "")
    if out_dir and os.path.isdir(out_dir):
        output_path = os.path.join(out_dir, os.path.basename(filename))
        try:
            with open(output_path, "wb") as fh:
                fh.write(file_bytes)
        except Exception:
            output_path = None

    elapsed = int((time.time() - t0) * 1000)
    out = {
        "status": "ok", "schema_version": "v1", "task_type": "excel_export",
        "elapsed_ms": elapsed,
        "summary": {
            "filename": os.path.basename(filename),
            "sheets": len(sheets),
            "sheet_names": [s["name"] for s in sheets],
            "total_rows": total_rows,
            "bytes": len(file_bytes),
            "output_path": output_path,
        },
        "summary_text": (
            f"✅ Excel 生成 · {os.path.basename(filename)} · "
            f"{len(sheets)} 表 · {total_rows} 行 · {len(file_bytes)//1024}KB"
            + (f"\n📁 {output_path}" if output_path else "")
        ),
    }
    if output_path:
        try:
            from _script_safety import write_output_artifact
            # 文件已落盘；helper 会以相同安全文件名覆盖为相同内容并返回统一 manifest。
            out["artifact_manifest"] = [
                write_output_artifact(os.path.basename(filename), file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            ]
        except Exception:
            out["artifact_manifest"] = [{
                "filename": os.path.basename(filename),
                "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "bytes": len(file_bytes),
                "local_path": output_path,
            }]
    # 小文件回 base64(无 OSS 也能取回)
    if return_base64 and len(file_bytes) <= 4 * 1024 * 1024:
        out["result_file_base64"] = base64.b64encode(file_bytes).decode("ascii")
        out["result_file_name"] = os.path.basename(filename)
    print(json.dumps(out, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
